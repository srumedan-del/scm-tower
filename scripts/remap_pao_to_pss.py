"""Create a safe copy of an outbound-detail workbook with PAO values remapped to PSS.

Example:
    python scripts/remap_pao_to_pss.py \
        C:\\Users\\its\\Downloads\\pssagustus.xlsx \
        pssagustus_pao-remapped.xlsx \
        PAO-2609-0416 PSS-2609-0759

For a new file, automatically remap every PAO to the preceding PSS with the
same project, branch, and location:
    python scripts/remap_pao_to_pss.py input.xlsx output.xlsx --auto-pao
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


CONTEXT_HEADERS = ("PROJECT", "CABANG/PERWAKILAN", "Location Code")


def normalize(value: object) -> str:
    return str(value or "").strip()


def is_document_type(value: object, prefix: str) -> bool:
    return normalize(value).upper().startswith(prefix)


def remap_document_number(
    source: Path, destination: Path, old_document_no: str, new_document_no: str
) -> int:
    workbook = load_workbook(source)
    changed_rows = 0

    for worksheet in workbook.worksheets:
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        document_column = headers.get("Document No.")
        if document_column is None:
            continue

        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=document_column)
            if str(cell.value or "").strip() == old_document_no:
                cell.value = new_document_no
                changed_rows += 1

    if changed_rows == 0:
        raise ValueError(f"Document number {old_document_no!r} was not found.")

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return changed_rows


def remap_all_pao_to_preceding_pss(source: Path, destination: Path) -> tuple[int, int]:
    """Map PAO rows to the preceding PSS in the same outbound context.

    If a PAO appears before its shipment row, the first following PSS is used.
    Rows without a PSS sharing their project, branch, and location are left
    unchanged, avoiding an assignment to an unrelated shipment.
    """
    workbook = load_workbook(source)
    changed_rows = 0
    skipped_rows = 0

    for worksheet in workbook.worksheets:
        headers = {
            normalize(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        document_column = headers.get("Document No.")
        if document_column is None:
            continue

        context_columns = [headers[header] for header in CONTEXT_HEADERS if header in headers]

        def context_for_row(row_number: int) -> tuple[str, ...]:
            return tuple(normalize(worksheet.cell(row_number, column).value) for column in context_columns)

        pss_by_context: dict[tuple[str, ...], list[tuple[int, str]]] = {}
        for row_number in range(2, worksheet.max_row + 1):
            document_no = worksheet.cell(row_number, document_column).value
            if is_document_type(document_no, "PSS-"):
                pss_by_context.setdefault(context_for_row(row_number), []).append(
                    (row_number, normalize(document_no))
                )

        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row_number, document_column)
            if not is_document_type(cell.value, "PAO-"):
                continue

            candidates = pss_by_context.get(context_for_row(row_number), [])
            if not candidates:
                skipped_rows += 1
                continue

            preceding = [candidate for candidate in candidates if candidate[0] <= row_number]
            if preceding:
                _, pss_no = max(preceding, key=lambda candidate: candidate[0])
            else:
                _, pss_no = min(candidates, key=lambda candidate: candidate[0])
            cell.value = pss_no
            changed_rows += 1

    if changed_rows == 0:
        raise ValueError("No PAO rows could be matched to a PSS in the same context.")

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return changed_rows, skipped_rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("old_document_no", nargs="?")
    parser.add_argument("new_document_no", nargs="?")
    parser.add_argument(
        "--auto-pao",
        action="store_true",
        help="Remap each PAO to the preceding PSS with the same project, branch, and location.",
    )
    args = parser.parse_args()

    if not args.source.is_file():
        parser.error(f"Source file was not found: {args.source}")

    if args.auto_pao:
        if args.old_document_no or args.new_document_no:
            parser.error("Do not provide document numbers together with --auto-pao.")
        changed_rows, skipped_rows = remap_all_pao_to_preceding_pss(args.source, args.destination)
        print(f"Remapped {changed_rows} PAO row(s) automatically.")
        if skipped_rows:
            print(f"Left {skipped_rows} PAO row(s) unchanged because no matching PSS was found.")
    else:
        if not args.old_document_no or not args.new_document_no:
            parser.error("Provide both document numbers, or use --auto-pao.")
        changed_rows = remap_document_number(
            args.source,
            args.destination,
            args.old_document_no,
            args.new_document_no,
        )
        print(f"Remapped {changed_rows} row(s) to {args.new_document_no}.")
    print(f"Created: {args.destination.resolve()}")


if __name__ == "__main__":
    main()
